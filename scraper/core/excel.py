"""Ekspor database paket ke Excel (.xlsx) dengan kolom sesuai kebutuhan pengguna."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (judul kolom, lebar)
KOLOM = [("NOMOR", 8), ("RUP", 12), ("NAMA PAKET", 62), ("JENIS KEGIATAN", 18), ("JALAN", 26), ("GANG", 34),
         ("KECAMATAN", 19), ("KELURAHAN", 19), ("MAK ASLI DARI PENGAMBILAN", 46), ("STATUS MAK", 21),
         ("ADA PERUBAHAN RUP", 30), ("PEMERIKSAAN", 52), ("PAGU", 19)]
KOLOM_TAHUN = ("TAHUN", 8)
_WARNA = {"Salah MAK": "F8D7D3", "MAK belum dipetakan": "FDF1C7", "Sesuai": "DDF1E4"}
_HEADER = PatternFill("solid", fgColor="1F3A5F")
_MERAH, _KUNING = PatternFill("solid", fgColor="F8D7D3"), PatternFill("solid", fgColor="FDF1C7")


def _gang(r):
    return "; ".join([f"Gg. {g}" for g in r.get("gang") or []] + [f"Komp. {k}" for k in r.get("komplek") or []])


def _mak_asli(r):
    """MAK lengkap seperti yang diambil dari SiRUP (tidak dipotong), tanpa pengulangan."""
    return "; ".join(dict.fromkeys(s.get("mak", "").strip() for s in r.get("sumber_dana_rincian") or [] if s.get("mak")))


def _perubahan_rup(r):
    if r.get("kode_rup_sebelumnya"):
        return f"Ya - revisi, menggantikan RUP {r['kode_rup_sebelumnya']}"
    if r.get("kode_rup_pengganti"):
        return f"Ya - sudah diganti RUP {r['kode_rup_pengganti']}"
    return "Tidak"


def _pemeriksaan(r):
    """'Kesalahan: A; B. Peringatan: C' atau 'OK'. Kedua nilai dikembalikan agar sel bisa diwarnai."""
    rinci = r.get("temuan_rinci") or []
    salah = list(dict.fromkeys(j for t, j, _ in rinci if t == "kesalahan"))
    peringatan = list(dict.fromkeys(j for t, j, _ in rinci if t == "peringatan"))
    if not salah and not peringatan:
        return "OK", None
    teks = ". ".join(x for x in (("Kesalahan: " + "; ".join(salah)) if salah else "",
                                 ("Peringatan: " + "; ".join(peringatan)) if peringatan else "") if x)
    return teks, ("kesalahan" if salah else "peringatan")


def _baris(r, nomor, dengan_tahun):
    pemeriksaan, tingkat = _pemeriksaan(r)
    isi = [nomor, r["kode_rup"], r["nama_paket"], r.get("jenis_kegiatan") or "",
           ("Jl. " + r["jalan"]) if r.get("jalan") else "", _gang(r), r.get("kecamatan") or "", r.get("kelurahan") or "",
           _mak_asli(r), r.get("status_mak") or "", _perubahan_rup(r), pemeriksaan, r.get("pagu") or 0]
    if dengan_tahun:
        isi.append(r["tahun"])
    return isi, pemeriksaan, tingkat


def _isi_sheet(ws, rows, dengan_tahun):
    kolom = KOLOM + ([KOLOM_TAHUN] if dengan_tahun else [])
    ws.append([k for k, _ in kolom])
    for i, (_, lebar) in enumerate(kolom, 1):
        ws.column_dimensions[get_column_letter(i)].width = lebar
        c = ws.cell(row=1, column=i)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    pagu_kol = len(KOLOM)
    for n, r in enumerate(rows, 1):
        isi, pemeriksaan, tingkat = _baris(r, n, dengan_tahun)
        ws.append(isi)
        baris = ws.max_row
        rup = ws.cell(row=baris, column=2)
        if r.get("link"):
            rup.hyperlink = r["link"]
            rup.font = Font(color="1F5FBF", underline="single")
        ws.cell(row=baris, column=pagu_kol).number_format = "#,##0"
        st = ws.cell(row=baris, column=10)
        if st.value in _WARNA:
            st.fill = PatternFill("solid", fgColor=_WARNA[st.value])
        if tingkat:
            ws.cell(row=baris, column=12).fill = _MERAH if tingkat == "kesalahan" else _KUNING
        if str(isi[10]).startswith("Ya"):
            ws.cell(row=baris, column=11).fill = _KUNING
    akhir = ws.max_row
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{max(akhir, 2)}"
    # total ikut berubah mengikuti filter Excel (SUBTOTAL 109 = jumlah baris yang terlihat)
    ws.cell(row=akhir + 2, column=pagu_kol - 1, value="TOTAL (sesuai filter)").font = Font(bold=True)
    tot = ws.cell(row=akhir + 2, column=pagu_kol, value=f"=SUBTOTAL(109,{get_column_letter(pagu_kol)}2:{get_column_letter(pagu_kol)}{max(akhir, 2)})")
    tot.font, tot.number_format = Font(bold=True), "#,##0"
    ws.cell(row=akhir + 2, column=pagu_kol - 1).alignment = Alignment(horizontal="right")


def buat_xlsx(per_tahun, gabungan=False):
    """per_tahun = {tahun: [baris database.baris()]}. Satu sheet per tahun (kolom persis sesuai kebutuhan);
    bila gabungan=True ada sheet 'Semua Tahun' paling depan dengan tambahan kolom TAHUN di ujung. -> bytes .xlsx"""
    wb = Workbook()
    wb.remove(wb.active)
    if gabungan:
        semua = [r for th in sorted(per_tahun) for r in per_tahun[th]]
        _isi_sheet(wb.create_sheet("Semua Tahun"), semua, True)
    for th in sorted(per_tahun):
        _isi_sheet(wb.create_sheet(str(th)), per_tahun[th], False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
