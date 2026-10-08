import http.client
import io
import tempfile
import threading
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

from scraper import web
from scraper.cli import muat_target
from scraper.core import database, db, excel, rekap

SAT = muat_target(None, None)[1]["id_satker"]
J26 = "1.04.05.2.01.0012.5.2.04.01.001.00004"
S26 = "1.04.05.2.01.0011.5.2.04.02.002.00004"
JUDUL = ["NOMOR", "RUP", "NAMA PAKET", "JENIS KEGIATAN", "JALAN", "GANG", "KECAMATAN", "KELURAHAN", "MAK ASLI DARI PENGAMBILAN",
         "STATUS MAK", "ADA PERUBAHAN RUP", "PEMERIKSAAN", "PAGU"]


def paket(kode, nama, pagu, th):
    return {"kode_rup": kode, "tahun": th, "klpd_nama": "K", "id_satker": SAT, "jenis": "penyedia", "nama_paket": nama,
            "penyelenggara": None, "pagu": pagu, "metode_pemilihan": "Tender", "sumber_dana": "APBD, APBD",
            "waktu_pemilihan": "Jan", "link": f"https://sirup.example/{kode}"}


def isi(conn, th, daftar):
    rid = db.mulai_run(conn, SAT, th)
    db.finalisasi(conn, rid, [paket(k, n, p, th) for k, n, p, _ in daftar], SAT, th)
    db.tutup_run(conn, rid, "success", len(daftar), len(daftar))
    rd = db.mulai_run(conn, SAT, th, "SIRUP_DETAIL")
    for k, n, p, m in daftar:
        db.simpan_detail(conn, rd, k, n, p, {"lokasi": [], "lokasi_ringkas": "", "volume": "1", "uraian": "Pekerjaan", "spesifikasi": "",
                                              "sumber_dana": [{"mak": m, "pagu": p, "sumber_dana": "APBD"}], "mak": m, "total_pagu": p, "extra": {}})


A = ("1", "Belanja Modal Jalan Kota (Jl. Flora, Gg. Flora 4, Kel. Siantan Hulu, Kec. Pontianak Utara)", 100, J26 + ".8.1.02.02.01.0001.00002")
B = ("2", "Belanja Modal Saluran Pembuang Pasang Surut (Jl. Flora, Komp. Indah, Kec. Pontianak Utara)", 40, J26)   # salah MAK
C = ("3", "Fotocopy", 10, "X.1.2")
A25 = ("10",) + A[1:]                      # kode RUP unik lintas tahun


def target(th):
    t = muat_target(None, th)[1]
    return t


class TestExcel(unittest.TestCase):
    def setUp(self):
        self.conn = db.buka(":memory:")
        self.addCleanup(self.conn.close)
        isi(self.conn, 2026, [A, B, C])
        t = target(2026)
        self.rows = database.baris(self.conn, t, rekap.lengkap(self.conn, t))

    def buka(self, isi_bytes):
        return load_workbook(io.BytesIO(isi_bytes))

    def test_kolom_persis_sesuai_kebutuhan_dan_isi_baris(self):
        wb = self.buka(excel.buat_xlsx({2026: self.rows}))
        ws = wb["2026"]
        self.assertEqual([c.value for c in ws[1]], JUDUL)
        baris = {r[1].value: [c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=4)}
        a = baris["1"]
        self.assertEqual((a[0], a[3], a[4], a[5], a[6], a[7]), (1, "Jalan", "Jl. Flora", "Gg. Flora 4", "Pontianak Utara", "Siantan Hulu"))
        self.assertEqual(a[8], A[3])                                      # MAK ASLI: tidak dipotong
        self.assertEqual((a[9], a[10], a[11], a[12]), ("Sesuai", "Tidak", "OK", 100))
        b = baris["2"]
        self.assertEqual((b[5], b[9]), ("Komp. Indah", "Salah MAK"))      # komplek ikut di kolom GANG
        self.assertTrue(b[11].startswith("Kesalahan: MAK tidak sesuai kategori"))

    def test_hyperlink_format_angka_filter_dan_total(self):
        ws = self.buka(excel.buat_xlsx({2026: self.rows}))["2026"]
        self.assertEqual(ws["B2"].hyperlink.target, "https://sirup.example/1")
        self.assertEqual(ws["M2"].number_format, "#,##0")
        self.assertEqual(ws.auto_filter.ref, "A1:M4")
        self.assertEqual(ws.freeze_panes, "D2")
        self.assertIn("SUBTOTAL(109,M2:M4)", ws["M6"].value)               # total mengikuti filter Excel

    def test_perubahan_rup_ya_dengan_nomor_lama_dan_baru(self):
        rows = [dict(r) for r in self.rows]
        rows[0]["kode_rup_sebelumnya"] = "111"
        rows[1]["kode_rup_pengganti"] = "222"
        ws = self.buka(excel.buat_xlsx({2026: rows}))["2026"]
        self.assertEqual(ws["K2"].value, "Ya - revisi, menggantikan RUP 111")
        self.assertEqual(ws["K3"].value, "Ya - sudah diganti RUP 222")
        self.assertEqual(ws["K4"].value, "Tidak")

    def test_semua_tahun_punya_sheet_gabungan_dengan_kolom_tahun_dan_sheet_per_tahun(self):
        r25 = [dict(self.rows[0], tahun=2025, kode_rup="9")]
        wb = self.buka(excel.buat_xlsx({2025: r25, 2026: self.rows}, gabungan=True))
        self.assertEqual(wb.sheetnames, ["Semua Tahun", "2025", "2026"])
        gab = wb["Semua Tahun"]
        self.assertEqual([c.value for c in gab[1]], JUDUL + ["TAHUN"])
        self.assertEqual([gab.cell(row=r, column=14).value for r in (2, 3, 4, 5)], [2025, 2026, 2026, 2026])
        self.assertEqual([c.value for c in wb["2026"][1]], JUDUL)           # sheet per tahun: kolom persis, tanpa TAHUN
        self.assertEqual([wb["2026"].cell(row=r, column=1).value for r in (2, 3, 4)], [1, 2, 3])   # nomor mulai 1 tiap sheet

    def test_kosong_tidak_error(self):
        wb = self.buka(excel.buat_xlsx({2026: []}))
        self.assertEqual([c.value for c in wb["2026"][1]], JUDUL)


class TestRuteExcel(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = self.path = str(Path(tmp.name) / "x.db")
        conn = db.buka(path)
        isi(conn, 2025, [A25])
        isi(conn, 2026, [A, B, C])
        conn.close()
        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), web.buat_handler(path, None))
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.addCleanup(self.srv.server_close)
        self.addCleanup(self.srv.shutdown)

    def get(self, jalur):
        c = http.client.HTTPConnection("127.0.0.1", self.srv.server_address[1], timeout=20)
        c.request("GET", jalur, headers={"Host": f"127.0.0.1:{self.srv.server_address[1]}"})
        r = c.getresponse()
        return r.status, dict(r.getheaders()), r.read()

    def test_per_tahun_dan_semua_tahun(self):
        s, h, isi_x = self.get("/api/excel?tahun=2026")
        self.assertEqual(s, 200)
        self.assertIn("spreadsheetml.sheet", h["Content-Type"])
        self.assertRegex(h["Content-Disposition"], r'attachment; filename="database_RUP_2026_\d{4}-\d{2}-\d{2}_\d{2}h\d{2}\.xlsx"')
        self.assertEqual(load_workbook(io.BytesIO(isi_x)).sheetnames, ["2026"])
        self.assertEqual(load_workbook(io.BytesIO(isi_x))["2026"].max_row, 3 + 1 + 2)       # header + 3 paket + baris total
        s, h, isi_x = self.get("/api/excel?tahun=semua")
        self.assertEqual((s, load_workbook(io.BytesIO(isi_x)).sheetnames), (200, ["Semua Tahun", "2025", "2026"]))
        self.assertIn("semua-tahun", h["Content-Disposition"])

    def test_paket_tidak_aktif_hanya_bila_diminta(self):
        conn = db.buka(self.path)
        isi(conn, 2026, [A, B])                      # pengambilan daftar berikutnya: paket 3 (Fotocopy) hilang -> tidak aktif
        conn.close()
        def jumlah(jalur):
            ws = load_workbook(io.BytesIO(self.get(jalur)[2]))["2026"]
            return [r[1].value for r in ws.iter_rows(min_row=2) if r[1].value and str(r[1].value).isdigit()]
        self.assertEqual(jumlah("/api/excel?tahun=2026"), ["1", "2"])
        self.assertEqual(jumlah("/api/excel?tahun=2026&tidak_aktif=1"), ["1", "2", "3"])

    def test_tahun_tanpa_data_ditolak_dengan_pesan(self):
        s, _, isi_x = self.get("/api/excel?tahun=1999")
        self.assertEqual(s, 404)


if __name__ == "__main__":
    unittest.main()
