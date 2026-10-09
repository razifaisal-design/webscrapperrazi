import http.client
import json
import tempfile
import threading
import time
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest import mock

from scraper import tugas, web
from scraper.core import db


class TestWebTugas(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = str(Path(self.tmp.name) / "t.db")
        db.buka(self.path).close()
        self.panggilan = []
        selesai = threading.Event()
        self.selesai = selesai

        def run_detail_palsu(conn, target, **kw):
            self.panggilan.append(("detail", target["tahun"], kw["koneksi"], kw["jeda"], kw["usia_hari"], kw["semua"]))
            kw["progres"](3, 10, 3, 0)
            kw["log"]("halo dari tugas")
            while not kw["berhenti"].is_set() and not selesai.is_set():
                time.sleep(0.01)
            return 130 if kw["berhenti"].is_set() else 0

        def run_daftar_palsu(conn, target, **kw):
            self.panggilan.append(("daftar", target["tahun"], kw["jeda"]))
            return 0

        def run_semua_palsu(conn, target, **kw):
            self.panggilan.append(("semua", target["tahun"], kw["koneksi"], kw["jeda"]))
            kw["tahap"](1, 2, "Daftar RUP")
            kw["log"]("tahap satu")
            kw["tahap"](2, 2, "Detail paket")
            kw["progres"](5, 20, 5, 0)
            return 0

        for nama, fn in (("run_detail", run_detail_palsu), ("run_daftar", run_daftar_palsu), ("run_semua", run_semua_palsu)):
            p = mock.patch.object(tugas, nama, fn)
            p.start()
            self.addCleanup(p.stop)
        self.peng = web.PengelolaTugas(self.path)
        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), web.buat_handler(self.path, None, self.peng))
        self.port = self.srv.server_address[1]
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.addCleanup(self.srv.server_close)
        self.addCleanup(self.srv.shutdown)
        self.addCleanup(self.selesai.set)

    def minta(self, metode, jalur, body=None, **header):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        h = {"Host": f"127.0.0.1:{self.port}", **header}
        data = None
        if body is not None:
            data = json.dumps(body)
            h["Content-Type"] = "application/json"
        c.request(metode, jalur, data, h)
        r = c.getresponse()
        isi = r.read()
        c.close()
        try:
            return r.status, json.loads(isi)
        except ValueError:
            return r.status, isi

    def mulai(self, **kw):
        return self.minta("POST", "/api/tugas/mulai", {"jenis": "detail", "tahun": 2026, "koneksi": 3, "jeda": 1, **kw}, **{"X-Pantau": "1"})

    # ---- keamanan ----
    def test_post_tanpa_header_dashboard_ditolak(self):
        self.assertEqual(self.minta("POST", "/api/tugas/mulai", {"jenis": "detail", "tahun": 2026})[0], 403)
        self.assertEqual(self.panggilan, [])

    def test_post_dari_situs_lain_ditolak(self):
        s, _ = self.minta("POST", "/api/tugas/mulai", {"jenis": "detail", "tahun": 2026}, **{"X-Pantau": "1", "Origin": "https://jahat.example"})
        self.assertEqual(s, 403)

    def test_host_palsu_ditolak(self):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", "/api/tugas", headers={"Host": "jahat.example"})
        self.assertEqual(c.getresponse().status, 403)
        c.close()

    def test_penanda_kode_basi_muncul_saat_kode_berubah_setelah_server_dimulai(self):
        self.assertFalse(self.minta("GET", "/api/tugas")[1]["kode_basi"])
        with mock.patch.object(web, "sidik_kode", return_value=web.sidik_kode() + 100):       # seolah ada berkas .py yang baru diubah
            self.assertTrue(self.minta("GET", "/api/tugas")[1]["kode_basi"])
        self.assertFalse(self.minta("GET", "/api/tugas")[1]["kode_basi"])

    def test_get_status_boleh(self):
        s, j = self.minta("GET", "/api/tugas")
        self.assertEqual((s, j["status"]), (200, "idle"))

    # ---- parameter ----
    def test_parameter_tidak_valid(self):
        for bad in ({"koneksi": 0}, {"koneksi": 11}, {"jeda": 0.05}, {"tahun": 1999}, {"usia_hari": -1}, {"jenis": "hapus"}):
            s, j = self.mulai(**bad)
            self.assertEqual(s, 400, bad)
        self.assertEqual(self.panggilan, [])

    # ---- alur ----
    def test_mulai_progres_henti(self):
        s, j = self.mulai(koneksi=3, jeda=0.7, usia_hari=3, semua=True)
        self.assertEqual((s, j["status"], j["koneksi"], j["jeda"]), (200, "berjalan", 3, 0.7))
        for _ in range(100):
            st = self.minta("GET", "/api/tugas")[1]
            if st["selesai"] == 3:
                break
            time.sleep(0.02)
        self.assertEqual((st["selesai"], st["total"], st["ok"]), (3, 10, 3))
        self.assertIn("halo dari tugas", st["log"])
        self.assertEqual(self.panggilan, [("detail", 2026, 3, 0.7, 3, True)])
        self.assertEqual(self.mulai()[0], 409)                          # tidak boleh dua tugas sekaligus
        self.minta("POST", "/api/tugas/henti", {}, **{"X-Pantau": "1"})
        for _ in range(100):
            st = self.minta("GET", "/api/tugas")[1]
            if st["status"] != "berjalan":
                break
            time.sleep(0.02)
        self.assertEqual((st["status"], st["kode"]), ("dihentikan", 130))

    def test_spse_nontender_tahun_angka_dan_semua(self):
        got = []

        def run_spse_palsu(conn, jenis, lpse, tahun, **kw):
            got.append((jenis, lpse, tahun, kw["jeda"]))
            return 0

        with mock.patch.object(tugas, "run_spse", run_spse_palsu):
            for th in (2026, "semua"):
                s, j = self.mulai(jenis="spse_nontender", tahun=th, koneksi=9, jeda=1.5)
                self.assertEqual((s, j["koneksi"]), (200, 1), th)                  # SPSE selalu 1 koneksi
                for _ in range(100):
                    if self.minta("GET", "/api/tugas")[1]["status"] != "berjalan":
                        break
                    time.sleep(0.02)
        self.assertEqual([(g[0], g[2]) for g in got], [("nontender", 2026), ("nontender", "semua")])
        self.assertEqual(self.mulai(jenis="spse_nontender", tahun=1999)[0], 400)
        self.assertEqual(self.mulai(jenis="spse_nontender", tahun=2026, jeda=0.05)[0], 400)

    def test_api_spse_bawaan_tahun_config(self):
        conn = db.buka(self.path)
        for kode, th in (("1", 2026), ("2", 2025)):
            conn.execute("INSERT INTO spse_paket (lpse,jenis,kode_paket,tahun,nama_paket,is_active) VALUES ('pontianak','nontender',?,?,?,1)",
                         (kode, th, f"Paket {kode}"))
        conn.commit()
        conn.close()
        s, j = self.minta("GET", "/api/spse")
        self.assertEqual(s, 200)
        self.assertEqual([b["kode_paket"] for b in j["baris"]], ["1"])
        self.assertEqual(sorted(j["tahun"]), [2025, 2026])
        s, j = self.minta("GET", "/api/spse?tahun=semua")
        self.assertEqual(len(j["baris"]), 2)

    def test_spse_semua_meneruskan_parameter(self):
        got = []

        def palsu(conn, jenis, lpse, tahun, **kw):
            got.append((jenis, lpse, tahun, kw["jeda"], kw["usia_hari"], kw["semua"], kw["satker"], kw["koneksi"]))
            kw["tahap"](1, 2, "Daftar paket SPSE")
            return 0

        def tunggu():
            for _ in range(100):
                if self.minta("GET", "/api/tugas")[1]["status"] != "berjalan":
                    return
                time.sleep(0.02)

        with mock.patch.object(tugas, "run_spse_semua", palsu):
            s, j = self.mulai(jenis="spse_semua", tahun=2026, koneksi=5, jeda=2, usia_hari=3, semua=True)
            self.assertEqual((s, j["koneksi"]), (200, 5))               # koneksi untuk tahap detail ikut dipakai
            tunggu()
            self.mulai(jenis="spse_semua", tahun="semua", koneksi=2, jeda=1.5, satker="DINAS X")
            tunggu()
            self.mulai(jenis="spse_semua", tahun=2026, jeda=1.5, satker="__tidak")
            tunggu()
        self.assertEqual(got[0][2:6], (2026, 2.0, 3, True))
        self.assertIsNone(got[0][6])                                  # tidak ada pilihan = semua satker (tidak dikunci ke satu dinas)
        self.assertEqual((got[0][7], got[1][2], got[1][6], got[1][7]), (5, "semua", "DINAS X", 2))
        self.assertEqual(got[2][6], "__tidak")                        # hanya Pengumuman

    def test_halaman_dan_api_spse_dan_banding(self):
        conn = db.buka(self.path)
        conn.execute("INSERT INTO spse_paket(lpse,jenis,kode_paket,tahun,nama_paket,tahapan,is_active,link) VALUES('pontianak','nontender','900',2026,'Paket Uji','Evaluasi Penawaran',1,'http://x/900')")
        conn.execute("INSERT INTO spse_detail(lpse,jenis,kode_paket,kode_rup,rup_json,satker,pagu,hps,lengkap,pemenang_terisi,kontrak_terisi,pemenang_nama,tahap) "
                     "VALUES('pontianak','nontender','900','55','[{\"kode_rup\":\"55\",\"nama_paket\":\"Paket Uji\",\"sumber_dana\":\"APBD\"}]',"
                     "'DINAS PERUMAHAN RAKYAT DAN KAWASAN PERMUKIMAN',1000,990,1,1,0,'CV Uji','Evaluasi Penawaran')")
        conn.execute("INSERT INTO spse_jadwal(lpse,jenis,kode_paket,no,tahap,mulai,sampai,jumlah_perubahan,riwayat_json) VALUES('pontianak','nontender','900',1,'Upload Dokumen Penawaran','2020-01-01T08:00','2020-01-05T08:00',0,'[]')")
        conn.execute("INSERT INTO sirup_paket(kode_rup,tahun,id_satker,jenis,nama_paket,pagu,metode_pemilihan,is_active,link) VALUES('55',2026,173394,'penyedia','Paket Uji',1000,'Pengadaan Langsung',1,'l')")
        conn.commit()
        conn.close()
        for jalur, kata in (("/spse", b"SPSE Non-Tender"), ("/banding", b"Perbandingan"), ("/gaya.css", b"--bg"), ("/bersama.js", b"pasangTugas"), ("/", b"gaya.css")):
            s, isi = self.minta("GET", jalur)
            self.assertEqual(s, 200, jalur)
            self.assertIn(kata, isi if isinstance(isi, bytes) else json.dumps(isi).encode(), jalur)
        s, j = self.minta("GET", "/api/spse")
        self.assertEqual(s, 200)
        b = [x for x in j["baris"] if x["kode_paket"] == "900"][0]
        self.assertEqual((b["lengkap"], b["upload_mulai"], b["kontrak_terisi"]), (1, "2020-01-01T08:00", 0))
        self.assertEqual((j["ringkas"]["terinci"], j["ringkas"]["pagu"]), (1, 1000))
        self.assertEqual([e["nama"] for e in j["satker_daftar"]], ["DINAS PERUMAHAN RAKYAT DAN KAWASAN PERMUKIMAN"])
        s, j = self.minta("GET", "/api/spse?satker=DINAS+TIDAK+ADA")                    # filter satker: tidak ada yang cocok
        self.assertEqual((j["ringkas"]["paket_aktif"], len(j["baris"])), (0, 0))
        s, j = self.minta("GET", "/api/spse?satker=dinas+perumahan+rakyat+dan+kawasan+permukiman")
        self.assertEqual(len(j["baris"]), 1)
        s, j = self.minta("GET", "/api/banding")
        self.assertEqual(s, 200)
        r = [x for x in j["baris"] if x["kode_rup"] == "55"][0]
        self.assertEqual((r["status"], r["kode_nontender"], r["pagu_sama"]), ("Sudah tayang", "900", True))
        s, j = self.minta("GET", "/api/spse/jadwal?kode=900")
        self.assertEqual((s, len(j["jadwal"])), (200, 1))
        self.assertEqual(self.minta("GET", "/api/spse/jadwal?kode=abc")[0], 400)
        for jalur in ("/api/spse/excel", "/api/banding/excel"):
            c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
            c.request("GET", jalur, headers={"Host": f"127.0.0.1:{self.port}"})
            r = c.getresponse()
            isi = r.read()
            c.close()
            self.assertEqual(r.status, 200, jalur)
            self.assertTrue(isi.startswith(b"PK"), jalur)             # berkas .xlsx = zip
            if jalur == "/api/banding/excel":                         # tautan harus di kolom yang benar: RUP -> SiRUP, non tender -> SPSE
                import io
                from openpyxl import load_workbook
                ws = load_workbook(io.BytesIO(isi)).active
                kepala = [c.value for c in ws[1]]
                baris = {c.value: c for c in ws[2]}
                self.assertEqual(kepala[3:5], ["KODE RUP", "KODE NON TENDER"])
                self.assertEqual(ws.cell(row=2, column=4).value, "55")
                self.assertEqual(ws.cell(row=2, column=4).hyperlink.target, "l")                # link SiRUP
                self.assertEqual(ws.cell(row=2, column=5).value, "900")
                self.assertEqual(ws.cell(row=2, column=5).hyperlink.target, "http://x/900")     # link SPSE

    def test_satker_baru_lewat_nama_rekap_dan_tugas_per_satker(self):
        import shutil
        from scraper import konfig
        cfg = Path(self.tmp.name) / "targets.json"
        shutil.copy(konfig.TARGETS, cfg)
        for modul in (konfig, web):
            p = mock.patch.object(modul, "TARGETS", cfg)
            p.start()
            self.addCleanup(p.stop)
        kunci = konfig.tambah_target("DINAS PEKERJAAN UMUM DAN PENATAAN RUANG", "Kota Pontianak", "D199", 173393)
        conn = db.buka(self.path)
        conn.execute("INSERT INTO sirup_paket(kode_rup,tahun,id_satker,jenis,nama_paket,pagu,metode_pemilihan,is_active,link) VALUES('777',2026,173393,'penyedia','Paket PU',500,'Pengadaan Langsung',1,'l')")
        conn.commit()
        conn.close()
        s, j = self.minta("GET", "/api/satker")
        self.assertEqual({x["satker"]: x["paket"] for x in j["satker"]}["DINAS PEKERJAAN UMUM DAN PENATAAN RUANG"], 1)
        s, j = self.minta("GET", f"/api/rekap?tahun=2026&target={kunci}")
        self.assertEqual((s, j["paket_aktif"], j["target"]["satker_nama"]), (200, 1, "DINAS PEKERJAAN UMUM DAN PENATAAN RUANG"))
        self.assertIsNone(j.get("klasifikasi"))                                # rekap umum: tanpa kategori Jalan/Saluran khusus Perkim
        s, j = self.minta("GET", f"/api/rekap?tahun=semua&target={kunci}")
        self.assertEqual((s, j["paket_dengan_detail"] + j["paket_aktif"]), (200, 1))
        s, j = self.minta("GET", "/api/rekap?tahun=2026")                      # satker bawaan tidak tercampur
        self.assertNotEqual(j["target"]["satker_nama"], "DINAS PEKERJAAN UMUM DAN PENATAAN RUANG")
        self.assertEqual([r["kode_rup"] for r in self.minta("GET", f"/api/database?tahun=2026&target={kunci}")[1]["baris"]], ["777"])
        self.assertEqual(self.minta("GET", f"/api/jalan?target={kunci}")[0], 200)
        self.assertEqual(self.minta("GET", f"/api/statistik?target={kunci}")[0], 200)
        self.assertEqual(self.mulai(target="tidak-ada")[0], 400)               # satker harus sudah terdaftar
        s, j = self.mulai(jenis="daftar", tahun=2026, target=kunci)
        self.assertEqual(s, 200)
        for _ in range(100):
            if self.panggilan:
                break
            time.sleep(0.02)
        self.assertEqual(self.panggilan[0][0], "daftar")

    def test_cari_dan_tambah_satker_lewat_nama(self):
        dipanggil = []

        def cari(nama, klpd=None, **kw):
            dipanggil.append(("cari", nama))
            return {"klpd": "Kota Pontianak", "tahun": 2026, "kandidat": [{"nama": "DINAS PERPUSTAKAAN DAN KEARSIPAN", "paket": 145, "cocok": "mengandung"}]}

        def tambah(nama, klpd=None, **kw):
            dipanggil.append(("tambah", nama))
            if nama == "ganda":
                raise ValueError("Nama 'ganda' cocok dengan 2 satker")
            return "dinas-perpustakaan", "DINAS PERPUSTAKAAN DAN KEARSIPAN"

        with mock.patch.object(tugas, "cari_satker_nama", cari), mock.patch.object(tugas, "tambah_satker_nama", tambah):
            s, j = self.minta("GET", "/api/satker/cari?q=perpustakaan")
            self.assertEqual((s, j["kandidat"][0]["nama"]), (200, "DINAS PERPUSTAKAAN DAN KEARSIPAN"))
            self.assertNotIn("id", j["kandidat"][0])                           # pengguna hanya melihat nama
            self.assertEqual(self.minta("POST", "/api/satker/tambah", {"nama": "x"})[0], 403)           # tanpa header dashboard ditolak
            s, j = self.minta("POST", "/api/satker/tambah", {"nama": "perpustakaan"}, **{"X-Pantau": "1"})
            self.assertEqual((s, j["satker_nama"]), (200, "DINAS PERPUSTAKAAN DAN KEARSIPAN"))
            s, j = self.minta("POST", "/api/satker/tambah", {"nama": "ganda"}, **{"X-Pantau": "1"})
            self.assertEqual(s, 400)
            self.assertIn("cocok dengan 2", j["error"])
        self.assertEqual(dipanggil, [("cari", "perpustakaan"), ("tambah", "perpustakaan"), ("tambah", "ganda")])

    def test_excel_tahan_karakter_kontrol_di_nama_paket(self):
        conn = db.buka(self.path)
        conn.execute("INSERT INTO spse_paket(lpse,jenis,kode_paket,tahun,nama_paket,tahapan,is_active,link) VALUES('pontianak','nontender','901',2026,?,'x',1,'l')", ("Paket\x0bkontrol",))
        conn.execute("INSERT INTO spse_detail(lpse,jenis,kode_paket,kode_rup,rup_json,satker,lengkap) VALUES('pontianak','nontender','901','1','[]',?,1)", ("DINAS\x0cX",))
        conn.commit()
        conn.close()
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", "/api/spse/excel", headers={"Host": f"127.0.0.1:{self.port}"})
        r = c.getresponse()
        isi = r.read()
        c.close()
        self.assertEqual(r.status, 200)
        self.assertTrue(isi.startswith(b"PK"))

    def test_tugas_periksa_rup_ke_sirup(self):
        got = []

        def kode(conn, tahun_list, satker):
            got.append(("kode", tahun_list, satker))
            return ["111", "222"]

        def periksa(conn, kodes, **kw):
            got.append(("periksa", kodes, kw["koneksi"], kw["jeda"], kw["ulang"]))
            return 0

        with mock.patch.object(tugas, "kode_rup_tak_berpasangan", kode), mock.patch.object(tugas, "run_periksa_rup", periksa):
            s, j = self.mulai(jenis="banding_periksa", tahun=2026, koneksi=3, jeda=1.5, semua=True, satker="DINAS X")
            self.assertEqual((s, j["koneksi"]), (200, 3))
            for _ in range(100):
                if self.minta("GET", "/api/tugas")[1]["status"] != "berjalan":
                    break
                time.sleep(0.02)
            self.mulai(jenis="banding_periksa", tahun="semua", koneksi=1, jeda=1.5, satker="__semua")
            for _ in range(100):
                if len(got) >= 4:
                    break
                time.sleep(0.02)
        self.assertEqual(got[0], ("kode", [2026], "DINAS X"))
        self.assertEqual(got[1], ("periksa", ["111", "222"], 3, 1.5, True))
        self.assertEqual(got[2], ("kode", None, None))                       # semua tahun, semua satker

    def test_id_satker_opsional_divalidasi_dan_diteruskan(self):
        for bad in ({"id_satker": "abc"}, {"id_satker": 0}, {"id_satker": -5}, {"id_satker": 10**10}):
            self.assertEqual(self.mulai(**bad)[0], 400, bad)
        self.assertEqual(self.panggilan, [])
        s, j = self.mulai(id_satker="69427", tahun=2021)                 # teks angka dari kolom input diterima
        self.assertEqual(s, 200)
        for _ in range(100):
            if self.panggilan:
                break
            time.sleep(0.02)
        self.assertEqual(self.panggilan[0][1], 2021)
        self.assertEqual(self.mulai(id_satker="")[0], 409)               # kosong = otomatis (valid), tapi tugas sebelumnya masih jalan

    def test_peringatan_laju_tinggi_dikirim_ke_ui(self):
        s, j = self.mulai(koneksi=10, jeda=0.2)
        self.assertIn("TINGGI", j["peringatan"])

    def test_satu_proses_semua_melaporkan_tahap(self):
        s, j = self.minta("POST", "/api/tugas/mulai", {"jenis": "semua", "tahun": 2025, "koneksi": 3, "jeda": 1.5}, **{"X-Pantau": "1"})
        self.assertEqual((s, j["jenis"], j["koneksi"]), (200, "semua", 3))
        for _ in range(100):
            st = self.minta("GET", "/api/tugas")[1]
            if st["status"] != "berjalan":
                break
            time.sleep(0.02)
        self.assertEqual((st["status"], st["tahap_ke"], st["tahap_total"], st["tahap_nama"]), ("selesai", 2, 2, "Detail paket"))
        self.assertEqual((st["selesai"], st["total"]), (5, 20))
        self.assertEqual(self.panggilan, [("semua", 2025, 3, 1.5)])

    def test_tugas_daftar_memakai_satu_koneksi(self):
        s, j = self.minta("POST", "/api/tugas/mulai", {"jenis": "daftar", "tahun": 2025, "koneksi": 5, "jeda": 2}, **{"X-Pantau": "1"})
        self.assertEqual((s, j["koneksi"]), (200, 1))
        for _ in range(100):
            if self.minta("GET", "/api/tugas")[1]["status"] != "berjalan":
                break
            time.sleep(0.02)
        self.assertEqual(self.panggilan, [("daftar", 2025, 2.0)])


class TestPerkiraanWaktu(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.p = web.PengelolaTugas(str(Path(self.tmp.name) / "x.db"))

    def progres(self, waktu, selesai, total=100):
        with mock.patch.object(web.time, "monotonic", return_value=waktu):
            self.p._progres(selesai, total, selesai, 0)
        return self.p.status()

    def test_kecepatan_dari_jendela_30_detik_dan_sisa_waktu(self):
        self.p._tahap_mulai = 0
        self.progres(10, 10)                                  # 1/detik sejak awal
        st = self.progres(20, 20)
        self.assertAlmostEqual(st["laju"], 1.0)
        self.assertEqual(st["sisa_detik"], 80)               # 80 tersisa / 1 per detik

    def test_laju_mengikuti_kecepatan_terbaru_bukan_rata_rata_seumur_hidup(self):
        self.p._tahap_mulai = 0
        self.progres(10, 1)                                   # awalnya lambat
        for t, n in ((100, 10), (110, 30)):                   # lalu 2 paket/detik di jendela terakhir
            st = self.progres(t, n)
        self.assertAlmostEqual(st["laju"], 2.0)
        self.assertEqual(st["sisa_detik"], 35)                # 70 / 2

    def test_selesai_sisa_nol_dan_belum_ada_data_kosong(self):
        self.p._tahap_mulai = 0
        self.assertIsNone(self.progres(5, 0)["sisa_detik"])
        self.assertEqual(self.progres(50, 100)["sisa_detik"], 0)

    def test_ganti_tahap_mereset_perkiraan(self):
        self.p._tahap_mulai = 0
        self.progres(10, 50)
        self.p._tahap(2, 2, "Detail paket")
        st = self.p.status()
        self.assertEqual((st["laju"], st["sisa_detik"], st["selesai"]), (None, None, 0))


if __name__ == "__main__":
    unittest.main()
